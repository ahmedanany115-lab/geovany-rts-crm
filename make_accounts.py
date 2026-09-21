import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "RTS ERP Accounts"

# All accounts
accounts = [
    # Email, Password, Full Name, Role, Department
    ("geovany.hany@rtegy.com",         "Geovany@153",  "Geovany Hany",          "Admin",        "IT"),
    ("dr.mohamed@rtegy.com",           "Ceo@123",      "Mohamed",               "Admin",        "Management"),
    ("Moataz@rtegy.com",               "Moataz@123",   "Moataz",                "Accountant",   "Finance"),
    ("fatma@rtegy.com",                "Fatma@123",    "Fatma",                 "Accountant",   "Finance"),
    ("ahmed.rekaby@rtegy.com",         "Rekaby@123",   "Ahmed Rekaby",          "Accountant",   "Finance"),
    ("Mrim@rtegy.com",                 "257993",       "Mrim",                  "Sales",        "Sales"),
    ("hanem.omar@rtegy.com",           "Hanem@123",    "Hanem Omar",            "Sales",        "Sales"),
    ("Abdelrahman.Abdullah@rtegy.com", "Abdo@123",     "Abdelrahman Abdullah",  "Sales",        "Sales"),
    ("Ahmed.khaled@rtegy.com",         "Akhaled@123",  "Ahmed Khaled",          "Sales",        "Sales"),
    ("Ahmed.Anany@rtegy.com",          "Anany@123",    "Ahmed Anany",           "SalesManager", "Sales"),
    ("khaled.taleb@rtegy.com",         "Taleb@123",    "Khaled Taleb",          "Purchasing",   "Purchasing"),
    ("Ahmed.Mostafa@rtegy.com",        "Mostafa@123",  "Ahmed Mostafa",         "SupportAgent", "Maintenance"),
    ("Hossam@rtegy.com",               "Hossam@123",   "Hossam",                "SupportAgent", "Maintenance"),
    ("Mostafa@rtegy.com",              "Mosta@123",    "Mostafa",               "SupportAgent", "Maintenance"),
    ("youssef.mounir@rtegy.com",       "Mounir@123",   "Youssef Mounir",        "SupportAgent", "Maintenance"),
    ("youssef.mohamed@rtegy.com",      "Youssef@123",  "Youssef Mohamed",       "SupportAgent", "Maintenance"),
    ("mahmoud.amr@rtegy.com",          "Mahmoud@123",  "Mahmoud Amr",           "SupportAgent", "Maintenance"),
    ("karim.mahmoud@rtegy.com",        "Karim@123",    "Karim Mahmoud",         "SupportAgent", "Maintenance"),
    ("randa@rtegy.com",                "Randa@123",    "Randa El Beheiry",      "SupportAgent", "Support"),
    ("mahmoud.nasrallah@rtegy.com",    "Nasrallah@123","Mahmoud Nasrallah",     "Delivery",     "Operations"),
    ("hany.mahmoud@rtegy.com",         "Hany@123",     "Hany Mahmoud",          "Delivery",     "Operations"),
    ("ahmed.reda@rtegy.com",           "Reda@123",     "Ahmed Reda",            "Delivery",     "Operations"),
    ("dina.reda@rtegy.com",            "Dina@123",     "Dina Reda",             "Marketing",    "Marketing"),
    ("mhy@rtegy.com",                  "Mahy@123",     "Mahy",                  "ReadOnly",     "—"),
    ("kfahim@rtegy.com",               "Kfahim@123",   "Kfahim",                "ReadOnly",     "—"),
    ("Dgeorge@rtegy.com",              "Daniel@123",   "Daniel George",         "ReadOnly",     "—"),
    ("Farah@rtegy.com",                "Farouha@123",  "Farah El Anany",        "ReadOnly",     "—"),
]

# Role colours
ROLE_FILLS = {
    "Admin":        "1E3A8A",  # dark blue
    "Accountant":   "065F46",  # dark green
    "SalesManager": "92400E",  # dark amber
    "Sales":        "B45309",  # amber
    "Purchasing":   "5B21B6",  # purple
    "SupportAgent": "0E7490",  # cyan
    "Delivery":     "374151",  # grey
    "Marketing":    "9D174D",  # pink
    "ReadOnly":     "4B5563",  # slate
}

def role_fill(role):
    hex_ = ROLE_FILLS.get(role, "374151")
    return PatternFill("solid", fgColor=hex_)

# Header row
headers = ["#", "Full Name", "Email", "Password", "Role", "Department"]
header_font  = Font(name="Arial", bold=True, color="FFFFFF", size=11)
header_fill  = PatternFill("solid", fgColor="1E293B")
header_align = Alignment(horizontal="center", vertical="center")

thin = Side(border_style="thin", color="D1D5DB")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

for col, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col, value=h)
    cell.font   = header_font
    cell.fill   = header_fill
    cell.alignment = header_align
    cell.border = border

ws.row_dimensions[1].height = 22

# Data rows
for row_idx, (email, pwd, name, role, dept) in enumerate(accounts, 2):
    row_data = [row_idx - 1, name, email, pwd, role, dept]
    is_odd = (row_idx % 2 == 0)
    row_bg  = "F8FAFC" if is_odd else "FFFFFF"

    for col, val in enumerate(row_data, 1):
        cell = ws.cell(row=row_idx, column=col, value=val)
        cell.font   = Font(name="Arial", size=10)
        cell.border = border
        cell.alignment = Alignment(vertical="center")

        if col == 5:  # Role column — coloured badge feel
            cell.fill  = role_fill(role)
            cell.font  = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        elif col == 1:  # #
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.fill = PatternFill("solid", fgColor=row_bg)
        else:
            cell.fill = PatternFill("solid", fgColor=row_bg)

    ws.row_dimensions[row_idx].height = 18

# Column widths
ws.column_dimensions["A"].width = 5
ws.column_dimensions["B"].width = 26
ws.column_dimensions["C"].width = 38
ws.column_dimensions["D"].width = 18
ws.column_dimensions["E"].width = 16
ws.column_dimensions["F"].width = 16

# Freeze header
ws.freeze_panes = "A2"

# Title above table
ws.insert_rows(1)
ws.merge_cells("A1:F1")
title_cell = ws["A1"]
title_cell.value = "RTS ERP — User Accounts"
title_cell.font  = Font(name="Arial", bold=True, size=14, color="1E293B")
title_cell.alignment = Alignment(horizontal="center", vertical="center")
title_cell.fill = PatternFill("solid", fgColor="E2E8F0")
ws.row_dimensions[1].height = 30

out = "/mnt/user-data/outputs/RTS-ERP-Accounts.xlsx"
wb.save(out)
print(f"Saved: {out}  ({len(accounts)} accounts)")
